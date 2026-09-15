from io import BytesIO

from openpyxl import load_workbook

from app.tests.helpers import auth_header, create_user_with_role


def test_admin_can_list_accounts(client, db_session):
    create_user_with_role(
        db_session,
        email="hr.list@test.com",
        password="hrpass123",
        role_name="hr",
        first_name="Amina",
        last_name="Benali",
    )

    response = client.get("/api/v1/users", headers=auth_header(client))
    assert response.status_code == 200
    data = response.json()
    emails = [user["email"] for user in data]
    assert "admin@test.com" in emails
    assert "hr.list@test.com" in emails

    hr = next(user for user in data if user["email"] == "hr.list@test.com")
    assert hr["full_name"] == "Amina Benali"
    assert [role["name"] for role in hr["roles"]] == ["hr"]
    assert "hashed_password" not in hr
    assert "password" not in hr


def test_hr_cannot_list_accounts(client, db_session):
    create_user_with_role(
        db_session,
        email="hr.list@test.com",
        password="hrpass123",
        role_name="hr",
    )
    response = client.get(
        "/api/v1/users",
        headers=auth_header(client, "hr.list@test.com", "hrpass123"),
    )
    assert response.status_code == 403


def test_unauthorized_cannot_list_accounts(client):
    response = client.get("/api/v1/users")
    assert response.status_code == 401


def test_admin_can_export_accounts_excel(client, db_session):
    create_user_with_role(
        db_session,
        email="hr.export@test.com",
        password="hrpass123",
        role_name="hr",
        first_name="Sara",
        last_name="Khelifi",
    )

    response = client.get("/api/v1/users/export", headers=auth_header(client))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "hr-platform-accounts.xlsx" in response.headers["content-disposition"]
    assert b"hashed_password" not in response.content
    assert b"hrpass123" not in response.content

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("ID", "First name", "Last name", "Email", "Roles", "Status", "Created at")
    emails = [row[3] for row in rows[1:]]
    assert "admin@test.com" in emails
    assert "hr.export@test.com" in emails

    hr_row = next(row for row in rows[1:] if row[3] == "hr.export@test.com")
    assert hr_row[1] == "Sara"
    assert hr_row[2] == "Khelifi"
    assert hr_row[4] == "hr"
    assert hr_row[5] == "Active"


def test_hr_cannot_export_accounts(client, db_session):
    create_user_with_role(
        db_session,
        email="hr.export@test.com",
        password="hrpass123",
        role_name="hr",
    )
    response = client.get(
        "/api/v1/users/export",
        headers=auth_header(client, "hr.export@test.com", "hrpass123"),
    )
    assert response.status_code == 403
